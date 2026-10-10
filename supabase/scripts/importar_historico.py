"""
ARGOS — importação do histórico da planilha Resultados_DUN_PROFESSIONAL_<ano>.xlsx

Lê a aba "--" (faturamento, rentabilidade e despesa de cada mês) e a aba
"Professional" (lista dos departamentos exclusivos) e gera um arquivo .sql
para rodar UMA VEZ no SQL Editor do Supabase, depois da migration inicial.

Os códigos dos departamentos não existem na planilha (ela casa pelo nome).
Para preencher os códigos, passe uma ou mais bases do Winthor em --bases.
Departamentos que ficarem sem código são criados assim mesmo; quando
aparecerem numa base futura, o ARGOS preenche o código pelo nome.

Uso:
    pip install openpyxl pandas xlrd
    python importar_historico.py Resultados_DUN_PROFESSIONAL_2026.xlsx \
        --bases 1464_prof_agosto.xls \
        --saida seed_historico_2026.sql

Opcional: --mapa mapa_codigos.csv  (colunas: nome;codigo) para completar códigos à mão.
"""

import argparse
import csv
import json
import os
import re
import sys
import uuid
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
import pandas as pd

MESES = {
    "JANEIRO": 1, "FEVEREIRO": 2, "MARÇO": 3, "MARCO": 3, "ABRIL": 4, "MAIO": 5,
    "JUNHO": 6, "JULHO": 7, "AGOSTO": 8, "SETEMBRO": 9, "OUTUBRO": 10,
    "NOVEMBRO": 11, "DEZEMBRO": 12,
}
SLUG = "dun-professional"
ORIGEM = "historico_planilha"


def norm(nome):
    return re.sub(r"\s+", " ", str(nome or "")).strip().upper()


def q(texto):
    """Literal SQL com aspas simples escapadas."""
    return "'" + str(texto).replace("'", "''") + "'"


def dec(valor, casas):
    return Decimal(repr(float(valor))).quantize(Decimal(1).scaleb(-casas), rounding=ROUND_HALF_UP)


def linha_total(ws):
    for r in range(1, ws.max_row + 1):
        if norm(ws.cell(r, 2).value) == "TOTAL":
            return r
    sys.exit(f"Não achei a linha TOTAL na aba {ws.title!r}.")


def ler_codigos(bases, mapa):
    codigos = {}
    for caminho in bases or []:
        df = pd.read_excel(caminho, header=None)
        for cod, nome in zip(df[0], df[1]):
            if pd.notna(cod) and pd.notna(nome):
                codigos[norm(nome)] = int(cod)
    if mapa:
        with open(mapa, encoding="utf-8-sig") as f:
            for row in csv.DictReader(f, delimiter=";"):
                if row.get("codigo"):
                    codigos[norm(row["nome"])] = int(row["codigo"])
    return codigos


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("planilha")
    ap.add_argument("--bases", nargs="*", default=[])
    ap.add_argument("--mapa")
    ap.add_argument("--saida", default="seed_historico.sql")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.planilha, data_only=True)
    ws = wb["--"]
    total_r = linha_total(ws)

    # ---- meses: cabeçalho na linha 2, blocos de 8 colunas a partir da coluna C
    meses = []
    for col in range(3, ws.max_column + 1, 8):
        titulo = norm(ws.cell(2, col).value)
        m = re.match(r"([A-ZÇ]+)-(\d{2})$", titulo)
        if not m or m.group(1) not in MESES:
            continue  # ANO, MÉDIA etc.
        ano, mes = 2000 + int(m.group(2)), MESES[m.group(1)]
        c_fat, c_rent, c_desp = col + 1, col + 2, col + 4
        linhas = []
        for r in range(4, total_r):
            nome = ws.cell(r, 2).value
            if not norm(nome):
                continue
            fat = ws.cell(r, c_fat).value or 0
            rent = ws.cell(r, c_rent).value or 0
            if fat == 0 and rent == 0:
                continue
            linhas.append((r, norm(nome), fat, rent))
        if not linhas:
            continue  # mês ainda não preenchido
        meses.append({
            "mes_ref": f"{ano:04d}-{mes:02d}-01",
            "linhas": linhas,
            "despesa": ws.cell(total_r, c_desp).value,
            "fat_planilha": ws.cell(total_r, c_fat).value,
            "rent_planilha": ws.cell(total_r, c_rent).value,
        })

    # ---- departamentos exclusivos (aba Professional)
    wp = wb["Professional"]
    tp = linha_total(wp)
    exclusivos = {norm(wp.cell(r, 2).value) for r in range(4, tp) if norm(wp.cell(r, 2).value)}

    # ---- códigos
    codigos = ler_codigos(args.bases, args.mapa)

    nomes = sorted({n for m in meses for (_, n, _, _) in m["linhas"]} | exclusivos)
    deptos = {n: {"id": str(uuid.uuid4()), "codigo": codigos.get(n)} for n in nomes}

    erros = []
    for m in meses:
        for (r, n, fat, rent) in m["linhas"]:
            if fat == 0 and rent != 0:
                erros.append(f"{m['mes_ref']} linha {r} {n}: rentabilidade sem faturamento")
        if m["despesa"] in (None, ""):
            erros.append(f"{m['mes_ref']}: despesa do mês vazia")
    if erros:
        sys.exit("Corrija antes de importar:\n  " + "\n  ".join(erros))

    # ---- SQL
    out = []
    w = out.append
    w("-- ARGOS — seed do histórico importado de " + os.path.basename(args.planilha))
    w("-- Rodar UMA VEZ no SQL Editor do Supabase, após a migration inicial.")
    w("begin;")
    w("do $$ begin")
    w(f"  if exists (select 1 from public.cargas_faturamento where origem = '{ORIGEM}') then")
    w("    raise exception 'ARGOS: o histórico da planilha já foi importado.';")
    w("  end if;")
    w("end $$;")
    w("")
    w("-- departamentos")
    for n in nomes:
        d = deptos[n]
        cod = "null" if d["codigo"] is None else str(d["codigo"])
        exc = "true" if n in exclusivos else "false"
        w("insert into public.departamentos (id, distribuidora_id, codigo_winthor, nome, exclusivo_professional) "
          f"select {q(d['id'])}, id, {cod}, {q(n)}, {exc} from public.distribuidoras where slug = {q(SLUG)};")

    for m in meses:
        carga = str(uuid.uuid4())
        tot_fat = sum(dec(f, 4) for (_, _, f, _) in m["linhas"])
        valores = []
        tot_rent = Decimal(0)
        for (r, n, fat, rent) in m["linhas"]:
            f4 = dec(fat, 4)
            pct = dec(rent / fat * 100, 10) if fat != 0 else Decimal(0)
            tot_rent += f4 * pct / 100
            original = json.dumps({"fonte": "planilha", "aba": "--", "linha": r,
                                   "fat_liq": fat, "rentab": rent}, ensure_ascii=False)
            cod = deptos[n]["codigo"]
            valores.append(f"({q(carga)}, {q(deptos[n]['id'])}, {'null' if cod is None else cod}, "
                           f"{q(n)}, {f4}, {pct}, {q(original)}::jsonb)")
        w("")
        w(f"-- {m['mes_ref']}  (planilha: FAT {m['fat_planilha']:.2f} | RENTAB {m['rent_planilha']:.2f} | DESPESA {m['despesa']})")
        w("insert into public.cargas_faturamento (id, distribuidora_id, mes_ref, versao, vigente, origem, "
          "arquivo_nome, qtd_linhas, total_faturamento, total_margem_valor, observacao) "
          f"select {q(carga)}, id, {q(m['mes_ref'])}, 1, true, {q(ORIGEM)}, {q(os.path.basename(args.planilha))}, "
          f"{len(valores)}, {tot_fat}, {tot_rent.quantize(Decimal('0.000001'))}, "
          f"'Importado da planilha (aba --)' from public.distribuidoras where slug = {q(SLUG)};")
        w("insert into public.faturamento_linhas (carga_id, departamento_id, codigo_winthor, "
          "nome_departamento, faturamento_liquido, margem_pct, linha_original) values")
        w(",\n".join("  " + v for v in valores) + ";")
        w("insert into public.lancamentos_despesa (distribuidora_id, mes_ref, versao, vigente, valor_total, origem, observacao) "
          f"select id, {q(m['mes_ref'])}, 1, true, {dec(m['despesa'], 2)}, {q(ORIGEM)}, "
          f"'Importado da planilha (aba --)' from public.distribuidoras where slug = {q(SLUG)};")

    w("")
    w("commit;")
    with open(args.saida, "w", encoding="utf-8") as f:
        f.write("\n".join(out) + "\n")

    sem_codigo = [n for n in nomes if deptos[n]["codigo"] is None]
    with open("departamentos_sem_codigo.csv", "w", encoding="utf-8", newline="") as f:
        wcsv = csv.writer(f, delimiter=";")
        wcsv.writerow(["nome", "codigo"])
        for n in sem_codigo:
            wcsv.writerow([n, ""])

    print(f"Meses importados: {', '.join(m['mes_ref'][:7] for m in meses)}")
    print(f"Departamentos: {len(nomes)} ({len(nomes) - len(sem_codigo)} com código, {len(sem_codigo)} sem)")
    print(f"Exclusivos PROFESSIONAL: {len(exclusivos)}")
    print(f"SQL gerado em: {args.saida}")
    if sem_codigo:
        print("Sem código listados em departamentos_sem_codigo.csv (opcional preencher e rodar com --mapa).")


if __name__ == "__main__":
    main()
